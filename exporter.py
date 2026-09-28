import os
import math
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config


def flatten_results(data_list):
    flattened = []
    for item in data_list:
        if isinstance(item, list):
            flattened.extend(item)
        else:
            flattened.append(item)
    return flattened


EXCEL_CELL_LIMIT = 32000


def _safe_value(value):
    """Excel-safe value: NaN->khali, ajeeb control characters hatao, lamba text kaato."""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (list, tuple, set, dict)):
        value = ", ".join(str(v) for v in value) if not isinstance(value, dict) else str(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > EXCEL_CELL_LIMIT:
            value = value[:EXCEL_CELL_LIMIT] + " ...[truncated]"
    return value


def export_to_structured_excel(data_list):
    records = flatten_results(data_list)
    if not records:
        raise ValueError("No data available to export.")

    df = pd.DataFrame(records)

    preferred = ["URL", "Source URL", "Status", "Error"]
    ordered = [column for column in preferred if column in df.columns]
    ordered += [column for column in df.columns if column not in ordered]
    df = df[ordered]

    wb = Workbook()
    ws = wb.active
    ws.title = "Scraped Data"

    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid",
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)

    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.append(list(df.columns))

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in df.itertuples(index=False, name=None):
        ws.append([_safe_value(v) for v in row])

    # Scraped text "=" se shuru ho to Excel usko FORMULA bana deta hai
    # (khatarnak ho sakta hai) — usko sirf text rakho
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_index in range(1, ws.max_column + 1):
        letter = get_column_letter(column_index)
        header = str(ws.cell(1, column_index).value or "")

        if "Content" in header or "Error" in header:
            width = 70
        elif "URL" in header or "Links" in header or "Image" in header:
            width = 45
        else:
            width = min(max(len(header) + 5, 15), 35)

        ws.column_dimensions[letter].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output = config.OUTPUT_FILENAME
    try:
        wb.save(output)
    except PermissionError:
        # File Excel mein khuli ho to Windows save nahi karne deta — naya naam
        base, ext = os.path.splitext(output)
        output = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
        wb.save(output)
        print("NOTE: Purani file Excel mein khuli thi, isliye naye naam se save ki.")
    print(f"SUCCESS: {output}")
    return output
