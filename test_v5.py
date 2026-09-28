"""V5 extra tests — real-world edge cases (py test_v5.py)."""
import os, tempfile, zipfile
from bs4 import BeautifulSoup
import openpyxl
import config, scraper, exporter
S = lambda h: BeautifulSoup(h, "html.parser")

# Blocked detection
assert not scraper._looks_blocked("Contact us", '<script src="https://www.google.com/recaptcha/api.js"></script>')
assert scraper._looks_blocked("Just a moment...", "<html>cf-chl-</html>")
assert scraper._looks_blocked("Please verify you are human", "")
assert not scraper._looks_blocked("access denied " + "article text " * 300, "")
# Contacts
c = scraper.scrape_contacts_and_links(S('<a href="mailto:a@b.com">Mail</a><a href="tel:+923001234567">Call</a><p>2026-09-25 10:30 id 123456789012345 logo@2x.png</p>'), "https://x.com/")
assert c["Emails Found"] == "a@b.com" and c["Phone Numbers"] == "+923001234567"
# Tables with duplicate headers
t = scraper.scrape_tables(S("<table><tr><th>Price</th><th>Price</th></tr><tr><td>1</td><td>2</td></tr></table>"), "u")
assert t[0]["Price"] == "1" and t[0]["Price_2"] == "2"
# Images skip base64
p = scraper.scrape_product_details(S('<img src="data:image/png;base64,AA"><img src="/r.jpg">'), "https://x.com/")
assert "data:" not in p["Image URLs"] and "r.jpg" in p["Image URLs"]
# Excel safety
fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd); old = config.OUTPUT_FILENAME; config.OUTPUT_FILENAME = path
exporter.export_to_structured_excel([{"URL": "a", "Content": "x\x0by" + "z" * 40000}, {"URL": "b", "Error": "=CMD()"}])
xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
assert "<f>" not in xml
ws = openpyxl.load_workbook(path).active
assert all(not isinstance(cl.value, float) for r in ws.iter_rows(min_row=2) for cl in r)
assert all(len(str(cl.value or "")) <= 32767 for r in ws.iter_rows() for cl in r)
os.remove(path); config.OUTPUT_FILENAME = old
print("V5 EDGE-CASE TESTS: PASS")
